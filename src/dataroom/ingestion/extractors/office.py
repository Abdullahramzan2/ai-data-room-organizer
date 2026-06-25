"""Microsoft Office document extractors."""

from __future__ import annotations

from pathlib import Path

from dataroom.ingestion.extractors.base import BaseExtractor, append_text_within_budget
from dataroom.ingestion.extractors.legacy_office import LegacyOfficeConverter
from dataroom.ingestion.models import ExtractedDocument, ExtractionMethod, FileMetadata


def read_docx_text(path: Path, max_chars: int) -> tuple[str, bool]:
    from docx import Document

    document = Document(str(path))
    text = ""
    truncated = False
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        text, para_truncated = append_text_within_budget(text, paragraph.text, max_chars)
        if para_truncated:
            truncated = True
            break
    return text, truncated


def read_pptx_text(path: Path, max_chars: int) -> tuple[str, int | None, bool]:
    from pptx import Presentation

    prs = Presentation(str(path))
    text = ""
    truncated = False
    for idx, slide in enumerate(prs.slides, start=1):
        text, slide_truncated = append_text_within_budget(text, f"## Slide {idx}", max_chars)
        if slide_truncated:
            truncated = True
            break
        for shape in slide.shapes:
            shape_text = getattr(shape, "text", "")
            if isinstance(shape_text, str) and shape_text.strip():
                text, shape_truncated = append_text_within_budget(text, shape_text, max_chars)
                if shape_truncated:
                    truncated = True
                    break
        if truncated:
            break
    return text, len(prs.slides), truncated


class DocxExtractor(BaseExtractor):
    extensions = {".docx"}

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            text, truncated = read_docx_text(path, self.max_text_chars)
            doc.text_content = text
            if truncated:
                doc.warnings.append(
                    f"Large document — only the first {self.max_text_chars:,} characters were loaded"
                )
        except Exception as exc:
            doc.errors.append(f"DOCX extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc


class DocExtractor(BaseExtractor):
    extensions = {".doc"}

    def __init__(
        self,
        ocr=None,
        max_text_chars: int = 500_000,
        legacy: LegacyOfficeConverter | None = None,
    ):
        super().__init__(ocr=ocr, max_text_chars=max_text_chars)
        self.legacy = legacy or LegacyOfficeConverter(ocr=ocr)

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            result = self.legacy.extract_doc(path, self.max_text_chars)
            doc.warnings.extend(result.warnings)
            doc.errors.extend(result.errors)

            if result.text.strip():
                doc.text_content = result.text
                doc.extra["legacy_extraction_method"] = result.method
                if result.method and result.method.endswith("_ocr"):
                    doc.ocr_applied = True
                    doc.ocr_text = result.text
                    doc.extraction_method = ExtractionMethod.OCR
                else:
                    doc.extraction_method = ExtractionMethod.NATIVE
            elif result.errors:
                doc.extraction_method = ExtractionMethod.FAILED
        except Exception as exc:
            doc.errors.append(f"DOC extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc


class XlsxExtractor(BaseExtractor):
    extensions = {".xlsx"}

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            lines: list[str] = []
            char_budget = self.max_text_chars
            for sheet in wb.worksheets:
                lines.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None and str(c).strip()]
                    if cells:
                        line = "\t".join(cells)
                        if len("\n".join(lines)) + len(line) > char_budget:
                            doc.warnings.append(
                                f"Large spreadsheet — only the first {char_budget:,} characters were loaded"
                            )
                            break
                        lines.append(line)
                if len("\n".join(lines)) >= char_budget:
                    break
            wb.close()
            doc.text_content = self._truncate("\n".join(lines))
        except Exception as exc:
            doc.errors.append(f"XLSX extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc


class XlsExtractor(BaseExtractor):
    extensions = {".xls"}

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            import xlrd

            book = xlrd.open_workbook(str(path))
            lines: list[str] = []
            char_budget = self.max_text_chars
            truncated = False
            for sheet in book.sheets():
                lines.append(f"## Sheet: {sheet.name}")
                for row_idx in range(sheet.nrows):
                    cells = [
                        str(sheet.cell_value(row_idx, col_idx))
                        for col_idx in range(sheet.ncols)
                        if str(sheet.cell_value(row_idx, col_idx)).strip()
                    ]
                    if cells:
                        line = "\t".join(cells)
                        if len("\n".join(lines)) + len(line) > char_budget:
                            doc.warnings.append(
                                f"Large spreadsheet — only the first {char_budget:,} characters were loaded"
                            )
                            truncated = True
                            break
                        lines.append(line)
                if truncated or len("\n".join(lines)) >= char_budget:
                    break
            doc.text_content = self._truncate("\n".join(lines))
        except Exception as exc:
            doc.errors.append(f"XLS extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc


class PptxExtractor(BaseExtractor):
    extensions = {".pptx"}

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            text, page_count, truncated = read_pptx_text(path, self.max_text_chars)
            doc.text_content = text
            doc.metadata.page_count = page_count
            if truncated:
                doc.warnings.append(
                    f"Large presentation — only the first {self.max_text_chars:,} characters were loaded"
                )
        except Exception as exc:
            doc.errors.append(f"PPTX extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc


class PptExtractor(BaseExtractor):
    extensions = {".ppt"}

    def __init__(
        self,
        ocr=None,
        max_text_chars: int = 500_000,
        legacy: LegacyOfficeConverter | None = None,
    ):
        super().__init__(ocr=ocr, max_text_chars=max_text_chars)
        self.legacy = legacy or LegacyOfficeConverter(ocr=ocr)

    def extract(self, path: Path, metadata: FileMetadata) -> ExtractedDocument:
        doc = ExtractedDocument(metadata=metadata, extraction_method=ExtractionMethod.NATIVE)
        try:
            result = self.legacy.extract_ppt(path, self.max_text_chars)
            doc.warnings.extend(result.warnings)
            doc.errors.extend(result.errors)

            if result.text.strip():
                doc.text_content = result.text
                doc.metadata.page_count = result.page_count
                doc.extra["legacy_extraction_method"] = result.method
                if result.method and result.method.endswith("_ocr"):
                    doc.ocr_applied = True
                    doc.ocr_text = result.text
                    doc.extraction_method = ExtractionMethod.OCR
                else:
                    doc.extraction_method = ExtractionMethod.NATIVE
            elif result.errors:
                doc.extraction_method = ExtractionMethod.FAILED
        except Exception as exc:
            doc.errors.append(f"PPT extraction failed: {exc}")
            doc.extraction_method = ExtractionMethod.FAILED
        return doc
