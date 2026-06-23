"""Write manifest.xlsx for classified documents."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dataroom.export.manifest import MANIFEST_COLUMNS


def _autosize_columns(ws: Worksheet, row_count: int) -> None:
    for col_idx, column_name in enumerate(MANIFEST_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(column_name)
        for row_idx in range(2, min(row_count + 2, 102)):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 60)


def write_manifest_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    """Write manifest rows to an Excel workbook with filter and frozen header."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Workbook has no active worksheet")
    ws.title = "Manifest"

    ws.append(list(MANIFEST_COLUMNS))
    for row in rows:
        ws.append([row.get(column, "") for column in MANIFEST_COLUMNS])

    ws.freeze_panes = "A2"
    if rows:
        last_col = get_column_letter(len(MANIFEST_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    _autosize_columns(ws, len(rows))
    wb.save(path)
