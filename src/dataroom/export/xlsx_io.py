"""Shared Excel read/write for tabular export artifacts."""

from __future__ import annotations

import math
import os
import tempfile
from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value)


def normalize_rows(rows: Sequence[Mapping[str, object]], columns: Sequence[str]) -> list[dict[str, str]]:
    return [{col: cell_str(row.get(col)) for col in columns} for row in rows]


def _autosize_columns(ws: Worksheet, columns: Sequence[str], row_count: int) -> None:
    for col_idx, column_name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(column_name)
        for row_idx in range(2, min(row_count + 2, 102)):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 60)


def _populate_sheet(ws: Worksheet, columns: Sequence[str], rows: list[dict[str, str]]) -> None:
    ws.append(list(columns))
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    ws.freeze_panes = "A2"
    if rows:
        last_col = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"
    _autosize_columns(ws, columns, len(rows))


def write_table_xlsx(
    path: Path,
    columns: Sequence[str],
    rows: Sequence[Mapping[str, object]],
    *,
    sheet_title: str = "Sheet1",
) -> None:
    """Write rows to an Excel workbook with frozen header and column filters."""
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized = normalize_rows(rows, columns)
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Workbook has no active worksheet")
    ws.title = sheet_title[:31]
    _populate_sheet(ws, columns, normalized)
    wb.save(path)


def write_table_xlsx_atomic(
    path: Path,
    columns: Sequence[str],
    rows: Sequence[Mapping[str, object]],
    *,
    sheet_title: str = "Sheet1",
) -> None:
    """Atomically replace an Excel file (temp file + os.replace)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(
        prefix=f"{path.stem}_",
        suffix=".tmp",
        dir=path.parent,
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        write_table_xlsx(tmp_path, columns, rows, sheet_title=sheet_title)
        os.replace(tmp_path, path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise


def read_table_xlsx(path: Path, columns: Sequence[str]) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            return []
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return []
        if not header:
            return []
        header_names = [cell_str(h) for h in header]
        col_map = {name: idx for idx, name in enumerate(header_names)}
        result: list[dict[str, str]] = []
        for values in rows_iter:
            if values is None or all(v is None for v in values):
                continue
            row_dict: dict[str, str] = {}
            for col in columns:
                idx = col_map.get(col)
                if idx is not None and idx < len(values):
                    row_dict[col] = cell_str(values[idx])
                else:
                    row_dict[col] = ""
            result.append(row_dict)
        return result
    finally:
        wb.close()
