"""Tests for dataroom rerun workflow."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import patch

import pytest
from click.testing import CliRunner

from dataroom.cli import main
from dataroom.export.review_queue import REVIEW_COLUMNS, write_review_queue_rows
from dataroom.pipeline import run_pipeline, run_rerun
from dataroom.pipeline.cache import load_classification_cache
from dataroom.pipeline.rerun import RerunError


ADMIN = "00_Admin_and_Index"


def test_review_queue_includes_corrected_folder_column(tmp_path: Path):
    input_dir = tmp_path / "in"
    input_dir.mkdir()
    (input_dir / "PSA_agreement.txt").write_text(
        "purchase and sale agreement PSA escrow closing conditions",
        encoding="utf-8",
    )
    output_dir = tmp_path / "out"
    run_pipeline(input_dir, output_dir, no_ocr=True)

    from openpyxl import load_workbook

    review_path = output_dir / ADMIN / "review_queue.xlsx"
    wb = load_workbook(review_path, read_only=True)
    try:
        ws = wb.active
        assert ws is not None
        header = [str(c or "") for c in next(ws.iter_rows(values_only=True))]
    finally:
        wb.close()
    assert header == REVIEW_COLUMNS
    assert "corrected_folder" in header


def test_run_rerun_applies_review_queue_correction(tmp_path: Path):
    input_dir = tmp_path / "in"
    input_dir.mkdir()
    source = input_dir / "BRAC_report.txt"
    source.write_text("brac remediation environmental", encoding="utf-8")
    output_dir = tmp_path / "out"

    run_pipeline(input_dir, output_dir, no_ocr=True)

    cls = load_classification_cache(output_dir / "classification_cache.json")
    original_folder = cls["results"][0]["category_folder"]
    corrected_folder = "01_Project_Overview"
    if original_folder == corrected_folder:
        corrected_folder = "04_Zoning_and_Land_Use"

    review_path = output_dir / ADMIN / "review_queue.xlsx"
    write_review_queue_rows(
        review_path,
        [
            {
                "file_name": source.name,
                "original_path": str(source),
                "assigned_folder": original_folder,
                "confidence": "low",
                "score": "0.3",
                "review_reason": "low confidence",
                "classification_reason": "test",
                "supporting_terms": "",
                "corrected_folder": corrected_folder,
            }
        ],
    )

    summary = run_rerun(output_dir)
    assert summary["rerun"] is True
    assert summary["corrections_applied"] == 1
    assert (output_dir / corrected_folder / source.name).is_file()

    cls_after = load_classification_cache(output_dir / "classification_cache.json")
    assert cls_after["results"][0]["category_folder"] == corrected_folder
    assert cls_after["results"][0]["needs_review"] is False

    run_summary = json.loads((output_dir / "run_summary.json").read_text(encoding="utf-8"))
    assert run_summary["rerun"] is True


def test_run_rerun_does_not_reingest(tmp_path: Path):
    input_dir = tmp_path / "in"
    input_dir.mkdir()
    source = input_dir / "doc.txt"
    source.write_text("purchase and sale agreement", encoding="utf-8")
    output_dir = tmp_path / "out"
    run_pipeline(input_dir, output_dir, no_ocr=True)

    with patch("dataroom.pipeline.run.run_ingestion_from_files") as mock_ingest:
        run_rerun(output_dir)
        mock_ingest.assert_not_called()


def test_run_rerun_missing_cache_raises(tmp_path: Path):
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    with pytest.raises(RerunError):
        run_rerun(output_dir)


@patch("dataroom.cli.run_rerun")
def test_rerun_cli(mock_run_rerun):
    mock_run_rerun.return_value = {
        "corrections_applied": 2,
        "organized": 5,
        "output_dir": "C:/out",
        "review_queue": "C:/out/00_Admin_and_Index/review_queue.xlsx",
        "review_queue_count": 1,
        "correction_warnings": [],
    }

    runner = CliRunner()
    with runner.isolated_filesystem():
        out = Path("data_room")
        out.mkdir()
        (out / "ingestion_cache.json").write_text("{}", encoding="utf-8")
        result = runner.invoke(main, ["rerun", str(out)])
    assert result.exit_code == 0
    assert "2 correction(s) applied" in result.output
