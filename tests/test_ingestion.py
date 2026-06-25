from pathlib import Path

from dataroom.ingestion.pipeline import run_ingestion
from dataroom.ingestion.router import ExtractionRouter
from dataroom.ingestion.extractors import build_extractors
from dataroom.ocr.tesseract import OcrConfig


def test_text_extraction(tmp_path: Path):
    sample = tmp_path / "notes.txt"
    sample.write_text("Project overview and site summary.", encoding="utf-8")

    router = ExtractionRouter(build_extractors(ocr=None))
    doc = router.extract(sample)

    assert "Project overview" in doc.combined_text
    assert doc.extraction_method.value == "native"
    assert doc.metadata.file_name == "notes.txt"


def test_pipeline_multiple_types(tmp_path: Path):
    (tmp_path / "readme.txt").write_text("Correspondence about wetlands.", encoding="utf-8")
    (tmp_path / "data.xlsx").write_bytes(_minimal_xlsx())

    result = run_ingestion(
        tmp_path,
        supported_extensions=[".txt", ".xlsx"],
        recursive=False,
        ocr_config=OcrConfig(enabled=False),
    )

    assert result.total_processed == 2
    assert result.total_with_text >= 1
    extensions = {d.metadata.extension for d in result.documents}
    assert ".txt" in extensions


def test_unsupported_extension_skipped(tmp_path: Path):
    (tmp_path / "binary.exe").write_bytes(b"\x00\x01")

    result = run_ingestion(
        tmp_path,
        supported_extensions=[".txt"],
        recursive=False,
        ocr_config=OcrConfig(enabled=False),
    )

    assert result.total_processed == 0
    assert len(result.skipped_files) == 0  # not even scanned as supported


def test_on_file_complete_called_per_file(tmp_path: Path):
    (tmp_path / "a.txt").write_text("first file", encoding="utf-8")
    (tmp_path / "b.txt").write_text("second file", encoding="utf-8")
    completed: list[tuple[str, str]] = []

    def on_file_complete(path: Path, outcome: str, error: str) -> None:
        completed.append((path.name, outcome))

    run_ingestion(
        tmp_path,
        supported_extensions=[".txt"],
        recursive=False,
        ocr_config=OcrConfig(enabled=False),
        on_file_complete=on_file_complete,
    )

    assert completed == [("a.txt", "ingested"), ("b.txt", "ingested")]


def _minimal_xlsx() -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    buf = BytesIO()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "USACE"
    ws["B1"] = "Section 404"
    wb.save(buf)
    return buf.getvalue()
