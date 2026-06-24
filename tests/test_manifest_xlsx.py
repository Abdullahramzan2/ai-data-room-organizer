"""Tests for manifest.xlsx export."""

from openpyxl import load_workbook

from dataroom.export import build_manifest_rows, write_manifest_xlsx
from dataroom.export.manifest import MANIFEST_COLUMNS


def test_write_manifest_xlsx(tmp_path):
    ingestion_docs = [
        {
            "source_path": "C:/docs/psa.txt",
            "file_name": "psa.txt",
            "extension": ".txt",
            "file_size": 1200,
            "modified_at": "2026-06-01T12:00:00+00:00",
            "file_hash": "abc123",
            "extraction_method": "native",
            "char_count": 100,
        }
    ]
    classification = [
        {
            "source_path": "C:/docs/psa.txt",
            "category_id": "02",
            "category_folder": "02_Land_Control_and_PSA",
            "confidence": "high",
            "score": 0.95,
            "method": "keyword",
            "reason": "Matched PSA",
            "supporting_terms": ["psa"],
            "entities": [],
            "needs_review": False,
            "api_used": False,
        }
    ]

    rows = build_manifest_rows(ingestion_docs, classification, tmp_path / "out")
    xlsx_path = tmp_path / "manifest.xlsx"
    write_manifest_xlsx(xlsx_path, rows)

    wb = load_workbook(xlsx_path)
    ws = wb["Manifest"]
    assert ws.cell(1, 1).value == "file_name"
    assert ws.cell(2, 1).value == "psa.txt"
    file_size_col = MANIFEST_COLUMNS.index("file_size") + 1
    modified_col = MANIFEST_COLUMNS.index("modified_at") + 1
    hash_col = MANIFEST_COLUMNS.index("file_hash") + 1
    assert ws.cell(2, file_size_col).value == "1200"
    assert ws.cell(2, modified_col).value == "2026-06-01T12:00:00+00:00"
    assert ws.cell(2, hash_col).value == "abc123"
    assert ws.freeze_panes == "A2"
    from openpyxl.utils import get_column_letter

    last_col = get_column_letter(len(MANIFEST_COLUMNS))
    assert ws.auto_filter.ref == f"A1:{last_col}2"
